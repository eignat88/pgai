import logging
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import psycopg2
import requests
from docx import Document

LOGGER = logging.getLogger(__name__)

SECTION_TYPES = {
    "Бизнес-процессы": "business_process",
    "Назначение, цель модификации": "purpose",
    "Изменение модели данных": "data_model",
    "Изменение интерфейса пользователей": "interface",
    "Алгоритмы": "algorithm",
    "Настройки": "settings",
    "Настройка": "settings",
    "Отчеты и выходные формы": "reports",
    "Размещение в системе": "placement",
    "Ограничение доступа": "access",
    "Настройка прав доступа": "access",
    "Допущения и ограничения": "limitations",
    "Связанные модификации": "related_modifications",
}

ENTITY_RE = re.compile(
    r"\b[A-Za-zА-Яа-я0-9_]*[A-Za-z]+[A-Za-z0-9_]*\.[A-Za-z0-9_]+"
    r"|\b[A-Z][A-Za-z0-9_]{2,}\b"
    r"|\btsd_[A-Za-z0-9_]+\b"
)
DAX_RE = re.compile(r"DAX-\d+")
VERSION_RE = re.compile(r"\b0\.\d+\b")
DATE_RE = re.compile(r"\b\d{2}\.\d{2}\.\d{4}\b")
AUTHOR_RE = re.compile(r"\b(Юрьева Ульяна|Игнатченко Евгений|Matveev Dmitriy)\b")

KNOWN_KEYWORDS = {
    "WMS_IsOversizedItemIM",
    "LFL_SCSPackTask",
    "WMSPickingRoute",
    "WMSOrderTrans",
    "PickingLineBuffer",
    "SalesTable",
    "InventTable",
    "WMSStoreArea",
    "WMS_TSDTaskType",
    "WMS_OperationType",
    "tsd_Setup",
    "tsd_AutoTaskTable",
    "tsd_WMSLocation",
}


@dataclass(slots=True)
class Settings:
    db_config: dict[str, str]
    docx_folder: Path
    ollama_url: str = "http://localhost:11434/api/embeddings"
    ollama_model: str = "nomic-embed-text"
    max_chunk_chars: int = 1800
    chunk_overlap: int = 250


def get_embedding(text: str, settings: Settings) -> list[float]:
    response = requests.post(
        settings.ollama_url,
        json={"model": settings.ollama_model, "prompt": text},
        timeout=120,
    )
    response.raise_for_status()
    body = response.json()
    if "embedding" not in body:
        raise ValueError("Ollama response does not contain 'embedding'")
    return body["embedding"]


def read_docx_text(path: Path) -> str:
    document = Document(path)
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)

    return "\n".join(parts)


def extract_section_text(text: str, section_title: str, max_chars: int = 1000) -> str | None:
    capture = False
    result: list[str] = []
    known_titles = set(SECTION_TYPES)

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if line == section_title:
            capture = True
            continue
        if capture and line in known_titles:
            break
        if capture and line:
            result.append(line)

    value = "\n".join(result).strip()
    return value[:max_chars] if value else None


def parse_doc_metadata(text: str, file_name: str) -> dict[str, Any]:
    dax = DAX_RE.search(file_name) or DAX_RE.search(text)

    first_non_empty_lines = [line.strip() for line in text.splitlines() if line.strip()]
    title = first_non_empty_lines[0] if first_non_empty_lines else file_name.rsplit(".", 1)[0]

    version_match = VERSION_RE.search(text)
    date_match = DATE_RE.search(text)
    author_match = AUTHOR_RE.search(text)

    document_date = None
    if date_match:
        document_date = datetime.strptime(date_match.group(0), "%d.%m.%Y").date().isoformat()

    return {
        "dax_code": dax.group(0) if dax else None,
        "title": title,
        "source_file": file_name,
        "version": version_match.group(0) if version_match else None,
        "business_process": extract_section_text(text, "Бизнес-процессы", max_chars=500),
        "purpose": extract_section_text(text, "Назначение, цель модификации", max_chars=1000),
        "author": author_match.group(1) if author_match else None,
        "document_date": document_date,
        "source_type": "docx",
    }


def split_into_sections(text: str) -> list[dict[str, str]]:
    sections: list[dict[str, str]] = []
    current_title = "Общее"
    current_type = "general"
    current_lines: list[str] = []

    def flush() -> None:
        if not current_lines:
            return
        sections.append(
            {
                "section_title": current_title,
                "chunk_type": current_type,
                "text": "\n".join(current_lines).strip(),
            }
        )

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line in SECTION_TYPES:
            flush()
            current_title = line
            current_type = SECTION_TYPES[line]
            current_lines = []
        else:
            current_lines.append(line)

    flush()
    return sections


def chunk_text(text: str, max_chars: int = 1800, overlap: int = 250) -> list[str]:
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    if len(text) <= max_chars:
        return [text]

    chunks: list[str] = []
    start = 0

    while start < len(text):
        end = min(start + max_chars, len(text))
        piece = text[start:end]

        cut = max(piece.rfind("\n"), piece.rfind(". "), piece.rfind("; "))
        if cut > 500:
            piece = piece[: cut + 1]
            end = start + cut + 1

        chunks.append(piece.strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)

    return [chunk for chunk in chunks if chunk]


def extract_entities(text: str) -> list[str]:
    found = {match.strip(".,;:()[]«»") for match in ENTITY_RE.findall(text)}
    found = {value for value in found if len(value) >= 3}
    found.update(keyword for keyword in KNOWN_KEYWORDS if keyword in text)
    return sorted(found)


def guess_process_type(text: str) -> str | None:
    lower = text.lower()
    if "упаков" in lower:
        return "packing"
    if "тсд" in lower or "tsd_" in text:
        return "tsd"
    if "сортиров" in lower or "scs" in lower:
        return "sorting"
    if "комплектац" in lower or "сборк" in lower:
        return "picking"
    if "номенклатур" in lower or "inventtable" in lower:
        return "item_master"
    return None


def insert_document(cur: Any, meta: dict[str, Any]) -> int:
    cur.execute(
        """
        INSERT INTO ai.fd_documents
            (title, source_file, dax_code, version, business_process, purpose, author, document_date, source_type)
        VALUES
            (%s, %s, %s, %s, %s, %s, %s, %s, %s)
        RETURNING id
        """,
        (
            meta["title"],
            meta["source_file"],
            meta["dax_code"],
            meta["version"],
            meta["business_process"],
            meta["purpose"],
            meta["author"],
            meta["document_date"],
            meta["source_type"],
        ),
    )
    return cur.fetchone()[0]


def insert_entity(cur: Any, name: str) -> int:
    cur.execute(
        """
        INSERT INTO ai.fd_entities (name)
        VALUES (%s)
        ON CONFLICT (name) DO UPDATE SET name = EXCLUDED.name
        RETURNING id
        """,
        (name,),
    )
    return cur.fetchone()[0]


def insert_chunk(
    cur: Any,
    settings: Settings,
    *,
    document_id: int,
    chunk: str,
    section_title: str,
    chunk_type: str,
    entities: list[str],
) -> int:
    embedding = get_embedding(chunk, settings)

    cur.execute(
        """
        INSERT INTO ai.fd_chunks
            (document_id, chunk_text, embedding, section_title, chunk_type, business_terms, jira_keys, process_type)
        VALUES
            (%s, %s, %s, %s, %s, %s, %s, %s)
        RETURNING id
        """,
        (
            document_id,
            chunk,
            embedding,
            section_title,
            chunk_type,
            [],
            DAX_RE.findall(chunk),
            guess_process_type(chunk),
        ),
    )
    chunk_id = cur.fetchone()[0]

    for entity in entities:
        entity_id = insert_entity(cur, entity)
        cur.execute(
            """
            INSERT INTO ai.fd_chunk_entities (chunk_id, entity_id)
            VALUES (%s, %s)
            ON CONFLICT DO NOTHING
            """,
            (chunk_id, entity_id),
        )

    return chunk_id


def insert_links_after_load(cur: Any) -> None:
    cur.execute("SELECT id, dax_code FROM ai.fd_documents WHERE dax_code IS NOT NULL")
    docs = {dax: doc_id for doc_id, dax in cur.fetchall()}

    cur.execute(
        """
        SELECT c.document_id, c.chunk_text
        FROM ai.fd_chunks c
        WHERE c.chunk_type = 'related_modifications'
        """
    )

    for source_document_id, text in cur.fetchall():
        for dax in set(DAX_RE.findall(text)):
            target_id = docs.get(dax)
            if not target_id or target_id == source_document_id:
                continue
            cur.execute(
                """
                INSERT INTO ai.fd_document_links
                    (source_document_id, target_document_id, link_type)
                VALUES
                    (%s, %s, %s)
                ON CONFLICT DO NOTHING
                """,
                (source_document_id, target_id, "related"),
            )


def ingest_docx_folder(settings: Settings) -> None:
    files = sorted({path.resolve() for path in settings.docx_folder.glob("*") if path.suffix.lower() == ".docx"})
    if not files:
        LOGGER.warning("Не найдены DOCX в папке: %s", settings.docx_folder)
        return

    with psycopg2.connect(**settings.db_config) as conn:
        with conn.cursor() as cur:
            for path in files:
                LOGGER.info("Загружаю: %s", path.name)
                text = read_docx_text(path)
                meta = parse_doc_metadata(text, path.name)
                document_id = insert_document(cur, meta)

                chunk_count = 0
                for section in split_into_sections(text):
                    chunks = chunk_text(
                        section["text"],
                        max_chars=settings.max_chunk_chars,
                        overlap=settings.chunk_overlap,
                    )
                    for chunk in chunks:
                        insert_chunk(
                            cur,
                            settings,
                            document_id=document_id,
                            chunk=chunk,
                            section_title=section["section_title"],
                            chunk_type=section["chunk_type"],
                            entities=extract_entities(chunk),
                        )
                        chunk_count += 1

                LOGGER.info("document_id=%s, chunks=%s", document_id, chunk_count)

            insert_links_after_load(cur)


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    settings = Settings(
        db_config={
            "dbname": "postgres",
            "user": "postgres",
            "password": "123",
            "host": "localhost",
            "port": "5432",
        },
        docx_folder=Path(r"C:\Users\ignatchenko\Documents\py\pg\docs"),
    )
    ingest_docx_folder(settings)
    LOGGER.info("Готово: документы загружены")


if __name__ == "__main__":
    main()
